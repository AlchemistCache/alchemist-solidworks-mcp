"""Production spreadsheets for sheet metal (openpyxl, no SolidWorks needed).

write_cutlist_xlsx() turns the sheet_metal_cutlist rows (C# SolidWorksService.SheetMetal.Features.cs) into a
formatted workbook: sheet «Ведомость» grouped by thickness + material, positions sorted by «Обозначение» (natural
order; identical designations with the same blank are merged and their quantities summed), per-group subtotals and a
grand total as live Excel formulas; sheet «Итоги» = the per-thickness totals."""
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# (header, row key | None for a formula, width, number format)
COLUMNS = [
    ("№", None, 5, "0"),
    ("Обозначение", "Обозначение", 22, "@"),
    ("Наименование", "Наименование", 26, "@"),
    ("Файл детали", "part", 26, "@"),
    ("Исполнение", "config", 14, "@"),
    ("Кол., шт", "qty", 8, "0"),
    ("Заготовка L, мм", "blank_length_mm", 11, "0.0"),
    ("Заготовка B, мм", "blank_width_mm", 11, "0.0"),
    ("Площадь нетто, м²", "net_area_m2", 11, "0.0000"),
    ("Длина реза, мм", "cut_length_mm", 11, "0.0"),
    ("Вырезы", "cut_outs", 8, "0"),
    ("Гибы", "bends", 7, "0"),
    ("R гиба, мм", "bend_radius_mm", 8, "0.0#"),
    ("K", "k_factor", 7, "0.00#"),
    ("Масса 1 шт, кг", "mass_kg", 10, "0.000"),
    ("Масса всего, кг", "=mass", 11, "0.000"),
    ("Площадь всего, м²", "=area", 11, "0.000"),
    ("Рез всего, м", "=cut", 10, "0.00"),
    ("Гибов всего", "=bends", 9, "0"),
]
_COL = {c[1]: get_column_letter(i + 1) for i, c in enumerate(COLUMNS) if c[1]}
_FORMULA = {"=mass": "={q}{r}*{mass_kg}{r}", "=area": "={q}{r}*{net_area_m2}{r}",
            "=cut": "={q}{r}*{cut_length_mm}{r}/1000", "=bends": "={q}{r}*{bends}{r}"}
_SUM_KEYS = ["qty", "=mass", "=area", "=cut", "=bends"]

_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD = PatternFill("solid", fgColor="D9D9D9")
_GROUP = PatternFill("solid", fgColor="EDF2F9")
_TOTAL = PatternFill("solid", fgColor="FFF2CC")


def _natural(s):
    return [int(x) if x.isdigit() else x.lower() for x in re.split(r"(\d+)", s or "")]


def _num(v):
    return round(v, 3) if isinstance(v, float) else v


def merge_rows(rows):
    """Same «Обозначение» + thickness + material + blank size = one position (qty summed, files listed)."""
    out, seen = [], {}
    for r in rows:
        des = (r.get("Обозначение") or "").strip()
        key = (des, r.get("thickness_mm"), r.get("material"), r.get("blank_length_mm"), r.get("blank_width_mm")) if des else None
        if key and key in seen:
            m = seen[key]
            m["qty"] = (m.get("qty") or 0) + (r.get("qty") or 0)
            if r.get("part") and r["part"] not in m["part"].split(", "):
                m["part"] += ", " + r["part"]
            continue
        r = dict(r)
        r["part"] = r.get("part") or ""
        out.append(r)
        if key:
            seen[key] = r
    return out


def _group_key(r):
    return (r.get("thickness_mm") or 0, r.get("material") or "?")


_HR = 3   # the header row of «Ведомость»
# sheet «Итоги»: (header, totals key, width, number format)
TOTAL_COLUMNS = [("Толщина | материал", "thickness_material", 30, "@"), ("Позиций", "positions", 9, "0"),
                 ("Кол., шт", "qty", 9, "0"), ("Площадь нетто, м²", "net_area_m2", 12, "0.000"),
                 ("Площадь заготовок, м²", "blank_area_m2", 12, "0.000"), ("Рез, м", "cut_length_m", 10, "0.00"),
                 ("Гибов", "bends", 9, "0"), ("Масса, кг", "mass_kg", 11, "0.000")]


def _head(cell, vertical="center"):
    """A header cell: bold on grey, bordered, centred and wrapped."""
    cell.font, cell.fill, cell.border = Font(bold=True), _HEAD, _BORDER
    cell.alignment = Alignment(horizontal="center", vertical=vertical, wrap_text=True)


def _fill_row(ws, r, fill):
    for j in range(1, len(COLUMNS) + 1):
        ws.cell(r, j).fill = fill


def _cutlist_header(ws):
    for i, (head, _, width, _) in enumerate(COLUMNS, 1):
        _head(ws.cell(_HR, i, head))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[_HR].height = 32
    ws.freeze_panes = ws.cell(_HR + 1, 3)


def _position_row(ws, r_i, pos, r):
    """One position: its number, the row's values and the per-position formulas (qty × …)."""
    refs = {k: _COL[k] for k in ("mass_kg", "net_area_m2", "cut_length_mm", "bends")}
    for j, (_, key, _, fmt) in enumerate(COLUMNS, 1):
        if key is None:
            v = pos
        elif key.startswith("="):
            v = _FORMULA[key].format(r=r_i, q=_COL["qty"], **refs)
        else:
            v = _num(r.get(key))
        cell = ws.cell(r_i, j, v)
        cell.number_format, cell.border = fmt, _BORDER


def _sum_cells(ws, r_i, formula):
    """Bold sums in the _SUM_KEYS columns of row r_i; formula(column letter) → the cell value."""
    for key in _SUM_KEYS:
        col = COLUMNS.index(next(c for c in COLUMNS if c[1] == key)) + 1
        cell = ws.cell(r_i, col, formula(get_column_letter(col)))
        cell.font, cell.number_format = Font(bold=True), COLUMNS[col - 1][3]


def _group_block(ws, r_i, pos, thk, mat, items):
    """One thickness + material group from row r_i: its title, the positions sorted by «Обозначение», the subtotal.
    Returns (the subtotal row, the last position number)."""
    items.sort(key=lambda r: (_natural(r.get("Обозначение")), _natural(r.get("part"))))
    ws.cell(r_i, 1, f"Лист {thk:g} мм — {mat}").font = Font(bold=True)
    _fill_row(ws, r_i, _GROUP)
    first = r_i + 1
    for r_i, r in enumerate(items, first):
        pos += 1
        _position_row(ws, r_i, pos, r)
    last = r_i
    sub = last + 1
    ws.cell(sub, 2, f"Итого {thk:g} мм — {mat}").font = Font(bold=True)
    _sum_cells(ws, sub, lambda L: f"=SUM({L}{first}:{L}{last})")
    for j in range(1, len(COLUMNS) + 1):
        ws.cell(sub, j).border = _BORDER
    return sub, pos


def _grand_total(ws, r_i, subtotal_rows):
    ws.cell(r_i, 2, "ВСЕГО").font = Font(bold=True, size=12)
    _sum_cells(ws, r_i, lambda L: "=" + "+".join(f"{L}{s}" for s in subtotal_rows) if subtotal_rows else 0)
    _fill_row(ws, r_i, _TOTAL)


def _print_setup(ws, last_row):
    """Filter over the table, landscape, one page wide, the header repeated on every page."""
    ws.auto_filter.ref = f"A{_HR}:{get_column_letter(len(COLUMNS))}{max(_HR, last_row)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{_HR}:{_HR}"


def _totals_sheet(ts, totals):
    for i, (h, _, w, _) in enumerate(TOTAL_COLUMNS, 1):
        _head(ts.cell(1, i, h), vertical=None)
        ts.column_dimensions[get_column_letter(i)].width = w
    for ri, t in enumerate(totals or [], 2):
        for i, (_, k, _, fmt) in enumerate(TOTAL_COLUMNS, 1):
            c = ts.cell(ri, i, _num(t.get(k)))
            c.number_format, c.border = fmt, _BORDER


def write_cutlist_xlsx(path: str, rows: list, totals: list, title: str = "") -> dict:
    rows = merge_rows(rows)
    groups = {}
    for r in rows:
        groups.setdefault(_group_key(r), []).append(r)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ведомость"
    ws.cell(1, 1, f"Ведомость раскроя листового металла{(' — ' + title) if title else ''}").font = Font(bold=True, size=13)
    _cutlist_header(ws)
    r_i, pos, subtotal_rows = _HR + 1, 0, []
    for (thk, mat), items in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        sub, pos = _group_block(ws, r_i, pos, thk, mat, items)
        subtotal_rows.append(sub)
        r_i = sub + 2
    _grand_total(ws, r_i, subtotal_rows)
    _print_setup(ws, r_i - 2)
    _totals_sheet(wb.create_sheet("Итоги"), totals)

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return {"xlsx": os.path.abspath(path), "positions": pos, "groups": len(groups)}


def write_table_xlsx(path: str, sheet: str, title: str, columns: list, rows: list, totals=None) -> dict:
    """Generic formatted table: columns = [(header, key, width)], rows = dicts; totals = (title, columns, rows)
    written below with a gap."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.cell(1, 1, title).font = Font(bold=True, size=13)

    def block(r0, cols, data):
        for i, (head, _, width) in enumerate(cols, 1):
            _head(ws.cell(r0, i, head))
            ws.column_dimensions[get_column_letter(i)].width = max(ws.column_dimensions[get_column_letter(i)].width or 0, width)
        for ri, row in enumerate(data, r0 + 1):
            for i, (_, key, _) in enumerate(cols, 1):
                c = ws.cell(ri, i, _num(row.get(key)))
                c.border = _BORDER
        return r0 + len(data) + 1

    r = block(3, columns, rows)
    ws.freeze_panes = ws.cell(4, 1)
    if totals and totals[2]:
        ws.cell(r + 1, 1, totals[0]).font = Font(bold=True)
        block(r + 2, totals[1], totals[2])
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return {"xlsx": os.path.abspath(path), "xlsx_rows": len(rows)}
