"""Specification (спецификация) per GOST 2.106-2019 form 1 from bom_export rows (one assembly level).

Sections in GOST order; a component's section is its «Раздел» property, else inferred: sub-assembly →
«Сборочные единицы», part with «Обозначение» → «Детали», part without it whose name cites a standard (ГОСТ, ОСТ,
DIN, ISO, EN …) → «Стандартные изделия», otherwise «Прочие изделия». Positions run through the sections (not in
«Документация»); assembly units and parts are sorted by «Обозначение», standard and other items by name."""
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from sheet_reports import _natural  # one natural sort for all reports

SECTIONS = ["Документация", "Комплексы", "Сборочные единицы", "Детали", "Стандартные изделия", "Прочие изделия",
            "Материалы", "Комплекты"]
_ALIASES = {s.lower(): s for s in SECTIONS}
_ALIASES.update({"сборочная единица": "Сборочные единицы", "деталь": "Детали", "стандартное изделие": "Стандартные изделия",
                 "прочее изделие": "Прочие изделия", "материал": "Материалы", "комплект": "Комплекты",
                 "документ": "Документация"})
_STANDARD = re.compile(r"\b(ГОСТ|ОСТ|DIN|ISO|EN|ANSI|ASME|JIS|ТУ)\b", re.I)
# GOST 2.106 form 1 column widths, mm → Excel character widths (~ mm / 2.1)
COLUMNS = [("Формат", 6), ("Зона", 6), ("Поз.", 8), ("Обозначение", 70), ("Наименование", 63), ("Кол.", 10),
           ("Примечание", 22)]


_PURCHASED = re.compile(r"[\\/](покупн|purchased|bought|vendor|модели производителей)", re.I)


def _inferred(row: dict) -> str:
    if (row.get("type") or "").lower() == "assembly":
        return "Сборочные единицы"
    if (row.get("Обозначение") or "").strip():
        return "Детали"
    name = " ".join(str(row.get(k) or "") for k in ("Наименование", "file"))
    return "Стандартные изделия" if _STANDARD.search(name) else "Прочие изделия"


def section_of(row: dict, moved: list = None) -> str:
    """The «Раздел» property, else inferred. «Детали» on a component without «Обозначение» that cites a standard or lies
    in a purchased-items folder is a template default (seen: every part, DIN nuts included, said «Детали») — inferred
    instead and noted in `moved`."""
    sec = (row.get("Раздел") or "").strip()
    if sec:
        sec = _ALIASES.get(sec.lower(), sec)
        if sec == "Детали" and not (row.get("Обозначение") or "").strip() and (row.get("type") or "").lower() != "assembly":
            name = " ".join(str(row.get(k) or "") for k in ("Наименование", "file"))
            why = ("the name cites a standard" if _STANDARD.search(name) else
                   "it lies in a purchased-items folder" if _PURCHASED.search(row.get("path") or "") else None)
            if why:
                new = _inferred(row)
                if moved is not None:
                    moved.append({"name": (row.get("Наименование") or row.get("file") or "").strip(), "to": new, "why": why})
                return new
        return sec
    return _inferred(row)


def build_spec(rows: list, doc_props: dict = None, assembly_drawing: bool = True, moved: list = None) -> list:
    """-> [{"section": name, "items": [{pos, format, zone, designation, name, qty, note}]}] in GOST order; `moved` gets
    the components whose «Детали» was overridden (section_of)."""
    doc_props = doc_props or {}
    by = {}
    for r in rows:
        by.setdefault(section_of(r, moved), []).append(r)
    out, pos = [], 0
    if assembly_drawing and doc_props.get("Обозначение"):
        out.append({"section": "Документация", "items": [
            {"pos": "", "format": "", "zone": "", "designation": doc_props["Обозначение"].strip() + " СБ",
             "name": "Сборочный чертёж", "qty": "", "note": ""}]})
    for sec in SECTIONS + sorted(k for k in by if k not in SECTIONS):
        items = by.get(sec)
        if not items:
            continue
        if sec in ("Стандартные изделия", "Прочие изделия", "Материалы"):
            items.sort(key=lambda r: _natural(r.get("Наименование") or r.get("file")))
        else:
            items.sort(key=lambda r: (_natural(r.get("Обозначение")), _natural(r.get("file"))))
        lst = []
        for r in items:
            if sec != "Документация":
                pos += 1
            name = (r.get("Наименование") or os.path.splitext(r.get("file") or "")[0]).strip()
            lst.append({"pos": pos if sec != "Документация" else "", "format": r.get("Формат") or "",
                        "zone": r.get("Зона") or "", "designation": (r.get("Обозначение") or "").strip(),
                        "name": name, "qty": r.get("qty"), "note": r.get("Примечание") or ""})
        if sec == "Документация" and out and out[0]["section"] == "Документация":
            out[0]["items"] += lst
        else:
            out.append({"section": sec, "items": lst})
    return out


def write_xlsx(path: str, spec: list, doc_props: dict = None, title: str = "") -> dict:
    doc_props = doc_props or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Спецификация"
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.cell(1, 1, "Спецификация " + " ".join(x for x in (doc_props.get("Обозначение"), doc_props.get("Наименование")) if x)
            + (f" ({title})" if title and not doc_props else "")).font = Font(bold=True, size=12)
    hr = 3
    for i, (h, mm) in enumerate(COLUMNS, 1):
        c = ws.cell(hr, i, h)
        c.font, c.border = Font(bold=True), border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(4.5, mm / 2.1)
    ws.freeze_panes = ws.cell(hr + 1, 1)
    r = hr + 1

    def line(values=(), bold_center_name=None):
        nonlocal r
        for j in range(1, len(COLUMNS) + 1):
            c = ws.cell(r, j, values[j - 1] if j - 1 < len(values) else None)
            c.border = border
            c.alignment = Alignment(horizontal="center" if j in (1, 2, 3, 6) else "left", vertical="top", wrap_text=True)
        if bold_center_name:
            c = ws.cell(r, 5, bold_center_name)
            c.font = Font(bold=True, underline="single")
            c.alignment = Alignment(horizontal="center")
        r += 1

    count = 0
    for sec in spec:
        line()
        line(bold_center_name=sec["section"])
        line()
        for it in sec["items"]:
            line((it["format"], it["zone"], it["pos"], it["designation"], it["name"], it["qty"], it["note"]))
            count += 1
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return {"xlsx": os.path.abspath(path), "rows": count}
